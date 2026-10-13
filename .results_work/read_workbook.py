import json
import openpyxl
from pathlib import Path

path = Path('E:/笔记/vla学习/论文复现/TransPrune_实验记录.xlsx')
wb = openpyxl.load_workbook(path, data_only=False)
out = {}
for s in wb:
    out[s.title] = {
        'rows': s.max_row, 'columns': s.max_column,
        'values': [[c.value for c in row] for row in s],
        'merged': list(map(str, s.merged_cells.ranges)),
        'freeze': s.freeze_panes,
    }
(path.parent / '.results_work/workbook_before.json').write_text(json.dumps(out, ensure_ascii=False, default=str), encoding='utf-8')
print(json.dumps({k: {'rows': v['rows'], 'columns': v['columns']} for k, v in out.items()}, ensure_ascii=True))
