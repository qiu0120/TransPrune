import json
from pathlib import Path
import openpyxl

root=Path('E:/笔记/vla学习/论文复现')
support=root/'.results_work'
out=root/'outputs/01a114ee-3e9e-7201-9c69-54d36a5b0930/TransPrune_实验记录.xlsx'
original=openpyxl.load_workbook(support/'workbook_original.xlsx', data_only=False)
edited=openpyxl.load_workbook(out, data_only=False)
cached=openpyxl.load_workbook(out, data_only=True)
source=json.loads((support/'server_results.json').read_text(encoding='utf-8-sig'))
parsed=json.loads((support/'parsed_results.json').read_text(encoding='utf-8'))

allowed={
    '主实验': ['A4','A13','K6','K19','H7:M9'],
    'MME分项': ['D15:F28','G7:H9','A11','A32:A33'],
    '实验安排': ['D8:D10','G8:H10','B11:D11','G11:H11','D14','G14:H14','C36'],
}
allowed_cells={}
for name,ranges in allowed.items():
    allowed_cells[name]=set()
    for area in ranges:
        block=edited[name][area]
        if isinstance(block,openpyxl.cell.cell.Cell):
            allowed_cells[name].add(block.coordinate)
        else:
            allowed_cells[name].update(c.coordinate for row in block for c in row)

changes=[]
for before in original:
    after=edited[before.title]
    assert set(map(str,before.merged_cells.ranges))==set(map(str,after.merged_cells.ranges)), before.title+' merges'
    assert before.freeze_panes==after.freeze_panes, before.title+' freeze'
    before_dv=[(str(d.sqref),d.type,d.formula1,d.formula2) for d in before.data_validations.dataValidation]
    after_dv=[(str(d.sqref),d.type,d.formula1,d.formula2) for d in after.data_validations.dataValidation]
    assert sorted(before_dv)==sorted(after_dv), before.title+' validation'
    for row in before:
        for c in row:
            value=after[c.coordinate].value
            if value != c.value:
                assert c.coordinate in allowed_cells.get(before.title,set()), (before.title,c.coordinate,c.value,value)
                changes.append((before.title,c.coordinate))

cols={'mme':'F','seed':'H','textvqa':'I','scienceqa':'J','pope':'K','gqa':'L','mmbench':'M'}
configs=['baseline','high','low']
recorded=0
for item in parsed['proof']:
    row=7+configs.index(item['config'])
    if item['benchmark']=='vqav2':
        assert cached['主实验'][f'G{row}'].value is None
    else:
        actual=cached['主实验'][cols[item['benchmark']]+str(row)].value
        if item['metric'] is None:
            assert actual is None
        else:
            assert abs(actual-item['metric'])<1e-9, (item,actual)
            recorded+=1
assert recorded==20
assert all(cached['主实验'][f'E{r}'].value is None for r in (7,8,9))
assert all(cached['主实验'][f'C{r}'].value is None for r in (7,8,9))
assert all(cached['效率'][f'D{r}'].value is None for r in (7,8,9))
assert len(edited['评测明细']['A7:A30'])==24
for s in cached:
    for row in s:
        for c in row:
            assert c.data_type!='e', (s.title,c.coordinate,c.value)

report={
    'verified_scores':recorded,
    'complete_prediction_runs':sum(r['complete'] for r in source['runs']),
    'pending_runs':[(r['benchmark'],r['config']) for r in source['runs'] if not r['complete']],
    'pending_online_scores':3,
    'changed_existing_cells':len(changes),
    'preserved_existing_tabs':original.sheetnames,
    'preserved_merges_freeze_panes_validations':True,
    'unrelated_values_and_formulas_unchanged':True,
    'formula_errors':0,
}
(support/'verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=True))
